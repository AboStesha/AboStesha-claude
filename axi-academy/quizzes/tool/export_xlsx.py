# Export questions.json as an Excel sheet for loading into the quiz system (one row per question, right-to-left).
# python3 tool/export_xlsx.py questions.json <out.xlsx>
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LETTERS = ['أ', 'ب', 'ج', 'د']
INK, MINT, STEEL, LINE = '090A0D', 'DFF8F0', 'EEF0F1', 'D9DDE0'
data = json.load(open(sys.argv[1], encoding='utf-8'))
plain = lambda s: s.replace('**', '')

wb = Workbook()
ws = wb.active
ws.title = 'بنك الأسئلة'
ws.sheet_view.rightToLeft = True
head = ['المعرّف', 'الكورس', 'المجموعة', 'رقم السؤال', 'السؤال', 'الخيار أ', 'الخيار ب', 'الخيار ج', 'الخيار د', 'الإجابة الصحيحة', 'نص الإجابة الصحيحة']
ws.append(head)
rows = []
for c in data['courses']:
    course = f"{c['label']} · {c['name']}"
    for block, group in [(l, f"الدرس {l['no']} · {l['title']}") for l in c['lessons']] + [(c['exam'], 'اختبار نهاية الكورس')]:
        for i, q in enumerate(block['questions'], 1):
            ws.append([q['id'], course, group, i, plain(q['q'])] + [plain(o) for o in q['options']] + [LETTERS[q['answer']], plain(q['options'][q['answer']])])
            rows.append(q['answer'])

thin = Side(style='thin', color=LINE)
for cell in ws[1]:
    cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    cell.fill = PatternFill('solid', fgColor=INK)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for r, answer in enumerate(rows, 2):
    for cell in ws[r]:
        cell.font = Font(name='Arial', size=11)
        cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
        cell.border = Border(bottom=thin)
    ws.cell(r, 6 + answer).fill = PatternFill('solid', fgColor=MINT)     # the correct option, in mint
    ws.cell(r, 10).font = Font(name='Arial', size=11, bold=True)
    ws.cell(r, 10).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(r, 4).alignment = Alignment(horizontal='center', vertical='top')
for col, width in zip('ABCDEFGHIJK', [14, 30, 34, 10, 60, 34, 34, 34, 34, 12, 40]):
    ws.column_dimensions[col].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

# summary: how the bank is meant to be used and how many questions each group has
sm = wb.create_sheet('ملخص')
sm.sheet_view.rightToLeft = True
sm.append(['الكورس', 'المجموعة', 'عدد الأسئلة'])
for c in data['courses']:
    for l in c['lessons']:
        sm.append([f"{c['label']} · {c['name']}", f"الدرس {l['no']} · {l['title']}", len(l['questions'])])
    sm.append([f"{c['label']} · {c['name']}", 'اختبار نهاية الكورس', len(c['exam']['questions'])])
sm.append([])
sm.append(['المجموع', '', len(rows)])
sm.append([])
for line in ['يظهر للطالب 5 أسئلة في كل اختبار، تُختار من مجموعة الدرس أو الكورس.',
             'إذا أخطأ الطالب في سؤال، يُستبدل بسؤال آخر من المجموعة نفسها لم يظهر له من قبل.',
             'الإجابة الصحيحة في العمود «الإجابة الصحيحة»، وخلية الخيار الصحيح مظلّلة.',
             'الأسعار والأرقام في الأسئلة أمثلة افتراضية للتعليم فقط.']:
    sm.append([line])
for cell in sm[1]:
    cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    cell.fill = PatternFill('solid', fgColor=INK)
for row in sm.iter_rows(min_row=2):
    for cell in row:
        cell.font = Font(name='Arial', size=11, bold=cell.row == sm.max_row - 5)
        cell.alignment = Alignment(horizontal='right', readingOrder=2)
for col, width in zip('ABC', [38, 46, 14]):
    sm.column_dimensions[col].width = width
wb.save(sys.argv[2])
print('wrote', sys.argv[2], len(rows), 'questions')
